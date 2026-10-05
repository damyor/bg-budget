#!/usr/bin/env python3
"""Extracts the named capital projects attached to the budget: the national priority
strategic investment projects (2026-2028 and the 2025 plan vs actual) and the municipal
investment programme (2025 payments, the 2026 list and the live status).

  python3 scripts/extract/capital_projects.py

Reads the cached downloads (see data/cache/projects/MANIFEST.md):
  data/cache/zdb-2026/zdb-2026.html               State Budget Act 2026, Annex 2 to art. 110
  data/cache/projects/report-2025-pr6.xlsx        2025 report, attachment pr.6 (priority projects)
  data/cache/zdb-2025/zdb-2025.html               State Budget Act 2025, Annex 2 (the plan pr.6 reports on)
  data/cache/municipalities/report-2025-pr7.xlsx  2025 report, attachment pr.7 (municipal programme)
  data/cache/projects/pms-103-2026-annex1.pdf     ПМС № 103/2026, Annex 1 (municipal projects 2026)
  data/cache/projects/ipop-projects.csv           ipop.mrrb.bg export (live status of the programme)

Writes data/sources/projects/:
  national-priority-projects-2026.csv         199 projects: capital spending 2026, forecasts 2027-2028
  national-priority-projects-2025-report.csv  pr.6 sections I (176) and II (228): 2025 plan vs actual
  municipal-investment-2025-report.csv        3,492 projects: transfer and Development Bank money paid in 2025
  municipal-investment-2026-decree.csv        3,492 projects: agreement, transfers to date, 2026 forecast, later
  municipal-investment-ipop.csv               3,492 projects: paid by year and by payer, as of the download

Checks, stopping on any failure: every table adds up to its printed total ("ОБЩО" rows of pr.6,
the summary row of pr.7, "Общо всичко:" of the decree, to the cent or the source's rounding);
pr.6's plan equals the State Budget Act 2025, Annex 2, project by project; project numbers run
1..n and codes are unique; on every pr.7 and ipop row the parts add up to the total; the three
municipal sources list the same 3,492 project codes, with the same municipality (ЕБК code in pr.7,
name in the decree and ipop, matched through data/sources/places/municipalities.csv) and, in pr.7
and ipop, the same project name. Needs bs4, lxml, openpyxl and pdfplumber.
"""

import bisect
import csv
import os
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal

import openpyxl
import pdfplumber
from bs4 import BeautifulSoup
from pdfplumber.utils import extract_text

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
CACHE = os.path.join(ROOT, 'data', 'cache')
OUT = os.path.join(ROOT, 'data', 'sources', 'projects')
REGISTER = os.path.join(ROOT, 'data', 'sources', 'places', 'municipalities.csv')

LOOKALIKES = dict(zip('ABCEHKMOPTXYaceopxy', 'АВСЕНКМОРТХУасеорху'))
NP_CODE = re.compile(r'^NP-\d\d\.\d{3}-\d{4}$')
OP_CODE = re.compile(r'^OP-\d\d\.\d{3}-\d{4}$')
# Province and municipality names the decree and ipop.mrrb.bg spell differently from the State Budget Act.
PROVINCE_ALIAS = {'софияобласт': 'софийска', 'софияград': 'софияград'}
MUNICIPALITY_ALIAS = {'столична': 'столичнаобщина'}


def fail(message):
    raise SystemExit(f'capital_projects.py: {message}')


def clean(text):
    """Whitespace collapsed, soft hyphens and non-breaking spaces removed."""
    return re.sub(r'\s+', ' ', str(text).replace('\xad', '').replace('\xa0', ' ')).strip()


def fix_cyrillic(text):
    """Latin look-alike letters inside Cyrillic words ("Oбезпечаване") → Cyrillic."""
    return re.sub(r'\w+', lambda m: ''.join(LOOKALIKES.get(c, c) for c in m.group(0)) if re.search('[Ѐ-ӿ]', m.group(0)) else m.group(0), text)


def norm(name):
    """A place name compared across sources: no quotes, spaces, hyphens or case."""
    return re.sub(r'[\s„“"”\'-]', '', fix_cyrillic(name)).lower()


def dec(value, places):
    """A spreadsheet number as an exact Decimal, rounded to the source's precision."""
    return Decimal(repr(value)).quantize(Decimal(1).scaleb(-places))


def fmt(value):
    return '' if value is None else str(value)


def write(name, header, rows):
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, name)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        w.writerows([[fmt(c) for c in r] for r in rows])
    print(f'  → {os.path.relpath(path, ROOT)} ({len(rows)} rows)')


# ---------- national priority projects 2026-2028 (State Budget Act 2026, Annex 2) ----------

def kilo(text):
    """'1 609,5' → Decimal('1609.5'); '–', '-' or blank → None. A trailing footnote mark is returned separately."""
    text = clean(text)
    mark = re.search(r'\*+$', text)
    text = text.rstrip('*').replace(' ', '')
    if text in ('', '–', '-'):
        return None, ''
    if not re.fullmatch(r'\d+(,\d+)?', text):
        fail(f'not an amount: {text!r}')
    return Decimal(text.replace(',', '.')), mark.group(0) if mark else ''


def national_2026():
    soup = BeautifulSoup(open(os.path.join(CACHE, 'zdb-2026', 'zdb-2026.html'), encoding='utf-8').read(), 'lxml')
    caption = soup.find(lambda t: t.name == 'p' and 'Приложение № 2 към чл. 110' in t.get_text())
    if caption is None:
        fail('Annex 2 to art. 110 not found in the State Budget Act 2026')
    table = caption.find_next('table')
    header = [clean(td.get_text()) for td in table.find('tr').find_all(['td', 'th'])]
    if header[3] != 'Размер на капиталови разходи по проекта през 2026 г.' or header[7] != 'Отговорна институция':
        fail(f'unexpected Annex 2 header: {header}')
    rows = []
    for tr in table.find_all('tr')[1:]:
        cells = [clean(td.get_text()) for td in tr.find_all(['td', 'th'])]
        if not any(cells):
            continue  # an empty row at a page break
        if len(cells) != 8 or not re.fullmatch(r'\d+', cells[0]) or not NP_CODE.match(cells[1]):
            fail(f'Annex 2: unexpected row {cells[:3]}')
        amounts = [kilo(c) for c in cells[3:6]]
        marks = ''.join(m for _, m in amounts)
        rows.append([int(cells[0]), cells[1], fix_cyrillic(cells[2]), *(a for a, _ in amounts), fix_cyrillic(cells[6]), fix_cyrillic(cells[7]), marks])
    notes = [clean(p.get_text()) for p in table.find_next_siblings('p')[:2]]
    if [r[0] for r in rows] != list(range(1, 200)) or len({r[1] for r in rows}) != 199:
        fail(f'Annex 2: {len(rows)} rows, expected projects 1-199 with unique codes')
    if not notes[1].startswith('** От които 2000 хил. евро') or [r[1] for r in rows if r[-1]] != ['NP-25.003-0019']:
        fail('Annex 2: the "**" footnote is not where it was')
    total = [sum(r[i] or 0 for r in rows) for i in (3, 4, 5)]
    by_institution = Counter()
    for r in rows:
        by_institution[r[7].split('/')[0]] += r[3] or 0
    print(f'Annex 2 (ЗДБ 2026): 199 projects; 2026 {total[0]:,} 2027 {total[1]:,} 2028 {total[2]:,} thousand EUR; '
          f'defence {by_institution["Министерство на отбраната"]:,}, regional development '
          f'{by_institution["Министерство на регионалното развитие и благоустройството"]:,}')
    write('national-priority-projects-2026.csv',
          ['no', 'code', 'name_bg', 'capex_2026_kEUR', 'capex_2027_kEUR', 'capex_2028_kEUR', 'indicator_bg', 'institution_bg', 'footnote'], rows)


# ---------- national priority projects 2025, plan vs actual (2025 report, attachment pr.6) ----------

def law_2025():
    """State Budget Act 2025, Annex 2 sections I and II: [(code, capital spending 2025 in thousand BGN)] in the printed order."""
    soup = BeautifulSoup(open(os.path.join(CACHE, 'zdb-2025', 'zdb-2025.html'), encoding='utf-8').read(), 'lxml')
    sections = {}
    for section, start in (('I', 'Приложение № 2, раздел I, към чл. 110, ал. 1'), ('II', 'Приложение № 2, раздел II, към чл. 110, ал. 2')):
        caption = soup.find(lambda t: t.name == 'p' and clean(t.get_text()) == start)
        if caption is None:
            fail(f'State Budget Act 2025: {start} not found')
        rows = []
        for tr in caption.find_next('table').find_all('tr'):
            cells = [clean(td.get_text()) for td in tr.find_all(['td', 'th'])]
            if len(cells) == 6 and re.fullmatch(r'\d+', cells[0]):
                rows.append((cells[1].rstrip('*'), kilo(cells[3])[0] or Decimal(0)))
        sections[section] = rows
    return sections


def national_2025():
    wb = openpyxl.load_workbook(os.path.join(CACHE, 'projects', 'report-2025-pr6.xlsx'), data_only=True)
    law = law_2025()
    out = []
    for section, ws in zip(('I', 'II'), wb.worksheets):
        sheet = list(ws.iter_rows(values_only=True))
        if not str(sheet[4][3]).startswith(('Размер на капиталови разходи', 'Индикативен размер')) or not str(sheet[4][6]).startswith('Отчет'):
            fail(f'pr.6 section {section}: unexpected header {sheet[4]}')
        printed = next(r for r in sheet if r[2] == 'ОБЩО')
        rows = []
        for r in sheet:
            if not isinstance(r[0], int):
                continue
            code = clean(r[1])
            mark = re.search(r'\*+$', code)
            code = code.rstrip('*')
            if not NP_CODE.match(code):
                fail(f'pr.6 section {section}: bad code {r[1]!r}')
            plan = None if r[3] in (None, '-') else dec(r[3], 3)
            actual = None if r[6] in (None, '-') else dec(r[6], 3)
            rows.append([section, r[0], code, mark.group(0) if mark else '', fix_cyrillic(clean(r[2])), plan, actual,
                         fix_cyrillic(clean(r[4] or '')), fix_cyrillic(clean(r[5]).replace('"', '„', 1).replace('"', '“', 1))])
        plan_sum = sum(r[5] or 0 for r in rows)
        actual_sum = sum(r[6] or 0 for r in rows)
        if [r[1] for r in rows] != list(range(1, len(rows) + 1)):
            fail(f'pr.6 section {section}: projects are not numbered 1..{len(rows)}')
        if abs(plan_sum - dec(printed[3], 3)) > Decimal('0.01') or abs(actual_sum - dec(printed[6], 3)) > Decimal('0.01'):
            fail(f'pr.6 section {section}: rows add up to {plan_sum} / {actual_sum}, "ОБЩО" is {printed[3]} / {printed[6]}')
        # The plan is the State Budget Act 2025's: same projects in the same order, same amounts (printed there to 0.1).
        if [code for code, _ in law[section]] != [r[2] for r in rows]:
            fail(f'pr.6 section {section}: the projects differ from the State Budget Act 2025, Annex 2')
        off = [r[2] for r, (_, amount) in zip(rows, law[section]) if abs((r[5] or 0) - amount) > Decimal('0.05')]
        if off:
            fail(f'pr.6 section {section}: plan differs from the State Budget Act 2025 for {off}')
        dup = [c for c, n in Counter(r[2] for r in rows).items() if n > 1]
        print(f'pr.6 section {section}: {len(rows)} projects, plan {plan_sum:,} actual {actual_sum:,} thousand BGN (= "ОБЩО"); '
              f'the plan equals the State Budget Act 2025, Annex 2 project by project (total {sum(a for _, a in law[section]):,})'
              + (f'; code printed twice: {", ".join(dup)}' if dup else ''))
        if section == 'I':
            own = [r for r in rows if not r[3]]
            print(f'  without the {len(rows) - len(own)} projects marked ** / *** (transfers to other budgets): {len(own)} projects, '
                  f'plan {sum(r[5] or 0 for r in own):,}, actual {sum(r[6] or 0 for r in own):,}')
        out += rows
    write('national-priority-projects-2025-report.csv',
          ['section', 'no', 'code', 'footnote', 'name_bg', 'plan_2025_kBGN', 'actual_2025_kBGN', 'indicator_bg', 'institution_bg'], out)


# ---------- municipal investment programme ----------

def read_register():
    with open(REGISTER, newline='', encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    by_name = {(norm(r['province']), norm(r['municipality'])): r['ebk_code'] for r in rows}
    names = {r['ebk_code']: r['municipality'] for r in rows}
    return by_name, names


def ebk_of(by_name, province, municipality):
    p, m = norm(province), norm(municipality)
    key = (PROVINCE_ALIAS.get(p, p), MUNICIPALITY_ALIAS.get(m, m))
    if key not in by_name:
        fail(f'no municipality {municipality!r} ({province}) in {os.path.relpath(REGISTER, ROOT)}')
    return by_name[key]


def municipal_2025(names):
    ws = openpyxl.load_workbook(os.path.join(CACHE, 'municipalities', 'report-2025-pr7.xlsx'), data_only=True).worksheets[0]
    sheet = list(ws.iter_rows(values_only=True))
    if sheet[2][:7] != ('ЕБК', 'Номер на проект', 'Приоритетен проект', 'Отговорна институция', 'Предоставен трансфер за 2025 г.',
                        'Предоставено финансиране за 2025 г.', 'Общо за 2025 г.'):
        fail(f'pr.7: unexpected header {sheet[2]}')
    printed = [dec(v, 5) for v in sheet[3][4:7]]
    rows = []
    for r in sheet[4:]:
        if r[1] is None:
            continue
        code, ebk = clean(r[1]), str(r[0])
        if not OP_CODE.match(code) or ebk not in names:
            fail(f'pr.7: bad row {r[:2]}')
        responsible = clean(r[3])
        # "Община Банско, област Благоевград"; Sofia is "Столична община, област София-град".
        named = re.sub(r', област .*$', '', responsible)
        if norm(named.removeprefix('Община ')) != norm(names[ebk]) and norm(named) != norm(names[ebk]):
            fail(f'pr.7 {code}: ЕБК {ebk} is {names[ebk]}, the row says {responsible}')
        transfer, bdb, total = (dec(v or 0, 5) for v in r[4:7])
        if abs(transfer + bdb - total) > Decimal('0.00002'):
            fail(f'pr.7 {code}: {transfer} + {bdb} != {total}')
        rows.append([ebk, code, clean(r[2]), responsible, transfer, bdb, total])
    sums = [sum(r[i] for r in rows) for i in (4, 5, 6)]
    if any(abs(s - p) > Decimal('0.0001') for s, p in zip(sums, printed)):
        fail(f'pr.7: rows add up to {sums}, the summary row is {printed}')
    if len(rows) != 3492 or len({r[1] for r in rows}) != 3492:
        fail(f'pr.7: {len(rows)} rows')
    paid = sum(1 for r in rows if r[6])
    print(f'pr.7: 3,492 projects in {len({r[0] for r in rows})} municipalities, {paid} with money in 2025; transfer {sums[0]:,.1f}, '
          f'Development Bank {sums[1]:,.1f}, total {sums[2]:,.1f} thousand BGN (= the summary row)')
    write('municipal-investment-2025-report.csv',
          ['ebk_code', 'code', 'name_bg', 'responsible_bg', 'transfer_2025_kBGN', 'bdb_2025_kBGN', 'total_2025_kBGN'], rows)
    return {r[1]: r for r in rows}


def euro(text):
    """'2 789 097,21 €' → Decimal; '- €' → 0; blank → None."""
    text = clean(text)
    if text == '':
        return None
    if text in ('- €', '-€'):
        return Decimal(0)
    m = re.fullmatch(r'(\d[\d ]*,\d\d) €', text)
    if not m:
        fail(f'decree: not an amount: {text!r}')
    return Decimal(m.group(1).replace(' ', '').replace(',', '.'))


def merge(values, tolerance=1.0):
    out = []
    for v in sorted(values):
        if not out or v - out[-1] > tolerance:
            out.append(v)
    return out


def decree_table(path):
    """The rows of the annex: every cell is the text of the characters inside its ruled borders (pages differ slightly)."""
    rows, totals, footnote = [], None, None
    with pdfplumber.open(path) as pdf:
        pages = len(pdf.pages)
        for pno, page in enumerate(pdf.pages, 1):
            cols = merge([(r['x0'] + r['x1']) / 2 for r in page.rects if r['width'] < 1.5 and r['height'] > 20])
            lines = merge([(r['top'] + r['bottom']) / 2 for r in page.rects if r['height'] < 1.5 and r['width'] > 30])
            if len(cols) != 11:
                fail(f'decree page {pno}: {len(cols)} column borders')
            cells, below = defaultdict(list), []
            for ch in page.chars:
                x, y = (ch['x0'] + ch['x1']) / 2, (ch['top'] + ch['bottom']) / 2
                if not cols[0] < x < cols[-1]:
                    continue  # the Gazette's running head
                if y > lines[-1]:
                    below.append(ch)
                    continue
                band, col = bisect.bisect(lines, y) - 1, bisect.bisect(cols, x) - 1
                if band >= 0:
                    cells[band, col].append(ch)
            if below:
                footnote = clean(extract_text(below))
            for band in range(len(lines) - 1):
                row = [clean(extract_text(cells.get((band, c), []))) for c in range(10)]
                if not any(row) or row[0].startswith('Номер на'):
                    continue
                if row[5] == 'Общо всичко:':
                    totals = [euro(c) for c in row[6:]]
                    continue
                if not OP_CODE.match(row[0]):
                    fail(f'decree page {pno}: a row without a project code: {row[:3]}')
                rows.append(row)
    return pages, rows, totals, footnote


def municipal_2026(by_name, pr7):
    pages, table, printed, footnote = decree_table(os.path.join(CACHE, 'projects', 'pms-103-2026-annex1.pdf'))
    if pages != 178 or len(table) != 3492 or printed is None:
        fail(f'decree: {pages} pages, {len(table)} rows, totals {printed}')
    if not (footnote or '').startswith('* Прогнозният бюджет за 2026 г. на проект не включва стойността на извършените плащания до 31.07.2026 г.'):
        fail(f'decree: footnote changed: {footnote}')
    rows = []
    for code, name, municipality, province, application, agreement, *amounts in table:
        # A dash the PDF puts before some municipality names ("- Кричим").
        municipality = municipality.removeprefix('- ')
        ebk = ebk_of(by_name, province, municipality)
        if code not in pr7 or pr7[code][0] != ebk:
            fail(f'decree {code}: {municipality} ({ebk}), pr.7 has {pr7.get(code, ["?"])[0]}')
        rows.append([code, name, municipality, province, ebk, application, '' if agreement == '0' else agreement, *(euro(a) for a in amounts)])
    sums = [sum(r[i] or 0 for r in rows) for i in (7, 8, 9, 10)]
    # The printed totals were summed from unrounded values: they differ from the rows by a few cents.
    if any(abs(s - p) > Decimal('0.10') for s, p in zip(sums, printed)):
        fail(f'decree: rows add up to {sums}, "Общо всичко:" is {printed}')
    if len({r[0] for r in rows}) != 3492 or {r[0] for r in rows} != set(pr7):
        fail('decree: the project codes differ from pr.7')
    print(f'ПМС № 103/2026 Annex 1: {pages} pages, 3,492 projects (the codes of pr.7, same municipalities); agreements {sums[0]:,}, '
          f'transfers {sums[1]:,}, 2026 forecast {sums[2]:,}, later {sums[3]:,} EUR (printed {", ".join(f"{p:,}" for p in printed)})')
    write('municipal-investment-2026-decree.csv',
          ['code', 'name_bg', 'municipality', 'province', 'ebk_code', 'application', 'agreement', 'agreement_EUR', 'transfers_EUR',
           'forecast_2026_EUR', 'later_EUR'], rows)


IPOP_HEADER = ['Област', 'Община', 'Проект OP', 'Описание', 'Стойност на споразумението между общината и МРРБ (EUR)',
               'Заявено в процес на проверка от МРРБ (EUR)', 'Одобрено за плащане от МРРБ и очакващо разплащане (EUR)',
               'Общо изплатено (EUR)', 'Изплатено от МРРБ (2024–2025) (EUR)', 'Изплатено от МРРБ (2026) (EUR)',
               'Изплатено от МРРБ — друга / неуточнена година (EUR)', 'Изплатено от ББР (EUR)']


def municipal_ipop(by_name, pr7):
    with open(os.path.join(CACHE, 'projects', 'ipop-projects.csv'), newline='', encoding='utf-8-sig') as f:
        header, *data = list(csv.reader(f, delimiter=';'))
    if header != IPOP_HEADER:
        fail(f'ipop: unexpected header {header}')
    rows = []
    for province, municipality, code, name, *amounts in data:
        values = [Decimal(a.replace(' ', '').replace(',', '.')) if a.strip() else Decimal(0) for a in amounts]
        agreement, verification, approved, paid, mrrb_2024_2025, mrrb_2026, mrrb_other, bdb = values
        if paid != mrrb_2024_2025 + mrrb_2026 + mrrb_other + bdb:
            fail(f'ipop {code}: paid {paid} is not the sum of its parts')
        ebk = ebk_of(by_name, province, municipality)
        if code not in pr7 or pr7[code][0] != ebk:
            fail(f'ipop {code}: {municipality} ({ebk}), pr.7 has {pr7.get(code, ["?"])[0]}')
        if clean(name) != pr7[code][2]:
            fail(f'ipop {code}: the name differs from pr.7')
        rows.append([code, province, municipality, ebk, *values])
    if len(rows) != 3492 or {r[0] for r in rows} != set(pr7):
        fail('ipop: the project codes differ from pr.7')
    sums = [sum(r[i] for r in rows) for i in range(4, 12)]
    print(f'ipop.mrrb.bg: 3,492 projects (the codes, municipalities and names of pr.7); agreements {sums[0]:,}, paid {sums[3]:,} '
          f'(regional development ministry 2024-2025 {sums[4]:,}, 2026 {sums[5]:,}, other {sums[6]:,}; Development Bank {sums[7]:,}), '
          f'claims under review {sums[1]:,}, approved and unpaid {sums[2]:,} EUR')
    write('municipal-investment-ipop.csv',
          ['code', 'province', 'municipality', 'ebk_code', 'agreement_EUR', 'in_verification_EUR', 'approved_unpaid_EUR', 'paid_EUR',
           'paid_mrrb_2024_2025_EUR', 'paid_mrrb_2026_EUR', 'paid_mrrb_other_EUR', 'paid_bdb_EUR'], rows)


def main():
    national_2026()
    national_2025()
    by_name, names = read_register()
    pr7 = municipal_2025(names)
    municipal_2026(by_name, pr7)
    municipal_ipop(by_name, pr7)


if __name__ == '__main__':
    sys.exit(main())
