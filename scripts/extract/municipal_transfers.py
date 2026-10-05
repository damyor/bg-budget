#!/usr/bin/env python3
"""Extracts what the central budget sends to each of the 265 municipalities, by
transfer type, from the State Budget Acts (State Gazette HTML), and a register of
the municipalities with their ЕБК codes and residents.

  python3 scripts/extract/municipal_transfers.py

Reads the cached downloads (see data/cache/*/MANIFEST.md):
  data/cache/zdb-<year>/zdb-<year>.html        State Budget Act 2024, 2025, 2026
  data/cache/municipalities/EBK-2026-public.xlsx  ЕБК 2026, section VII (municipality codes)
  data/cache/municipalities/report-2025-pr7.xlsx  2025 report, annex pr.7 (codes, cross-check)
  data/cache/municipalities/Pop_6.1.1_Pop_DR.xlsx NSI population by municipality at 31 Dec

Writes:
  data/sources/budget-<year>/municipal-transfers-<year>.csv  one row per municipality and
      transfer type (art. 53 in 2024 and 2025, art. 51 in 2026), plus the printed totals
  data/sources/places/municipalities.csv  ЕБК code, province, name, residents at 31 Dec 2023-2025

Checks, stopping on any failure: every column adds up to the act's "ВСИЧКО:" row and to
the amounts in the article's text; on every row column 2 = 3 + 4 + 5 + 6; the grand total
equals "Общините" in art. 1(2); the delegated-activities column equals, municipality by
municipality, the total of the art. 52/54 table, and that table equals the existing
municipal-delegated-<year>.csv cell by cell; every municipality has exactly one ЕБК code
(the ЕБК and annex pr.7 agree); the NSI municipalities add up to their provinces and to
the national total. Needs bs4, lxml and openpyxl.
"""

import csv
import os
import re
import sys

import openpyxl
from bs4 import BeautifulSoup

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
CACHE = os.path.join(ROOT, 'data', 'cache')
SOURCES = os.path.join(ROOT, 'data', 'sources')

YEARS = {
    # year: (transfers article, delegated-activities article, unit)
    2024: ('53', '54', 'kBGN'),
    2025: ('53', '54', 'kBGN'),
    2026: ('51', '52', 'kEUR'),
}
# Columns 3-7 of the transfers table: id, the words that start its amount in the article's text.
TRANSFERS = [
    ('delegated', 3, r'обща субсидия за делегираните от държавата дейности'),
    ('equalising', 4, r'обща изравнителна субсидия'),
    ('winter_roads', 5, r'снегопочистване на общински пътища'),
    ('capital', 6, r'целева субсидия за капиталови разходи'),
    ('other_targeted', 7, r'трансфери за други целеви разходи за местни дейности'),
]
DELEGATED_COLUMNS = ['total_delegated(1=2+3+4+5+6+7+8)', 'municipal_administration(2)', 'of_which_mayors(2a)', 'of_which_admin_staff(2b)',
                     'defence_security(3)', 'education(4)', 'health(5)', 'social_services(6)', 'culture(7)', 'economic_activities(8)']
SOFIA = 'Столична община'
LOOKALIKES = dict(zip('ABCEHKMOPTXYaceopxy', 'АВСЕНКМОРТХУасеорху'))
NSI_SHEETS = {2023: '2023', 2024: '2024', 2025: '2025'}  # population at 31 Dec of the year


def clean(text):
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()


def fix_cyrillic(text):
    """Latin look-alike letters inside Cyrillic words ("Cмолян") → Cyrillic."""
    return re.sub(r'\w+', lambda m: ''.join(LOOKALIKES.get(c, c) for c in m.group(0)) if re.search('[Ѐ-ӿ]', m.group(0)) else m.group(0), text)


def amount(text):
    """'11 681,1' → 11681.1; blank (or a dash) → None."""
    text = clean(text).replace(' ', '')
    if text in ('', '-', '–'):
        return None
    if not re.fullmatch(r'-?\d+(,\d+)?', text):
        raise ValueError(f'not an amount: {text!r}')
    return float(text.replace(',', '.'))


def r1(x):
    return round(x + 0.0, 1)


def province_of(heading):
    """'ОБЛАСТ ВЕЛИКО ТЪРНОВО' → 'Велико Търново' (the 2025 Gazette splits one heading as 'ОБ ЛАСТ')."""
    name = re.sub(r'^ОБ\s*ЛАСТ\s+', '', fix_cyrillic(heading).strip(), flags=re.I)
    return re.sub(r'(^|[\s-])(\w)', lambda m: m.group(1) + m.group(2).upper(), name.lower()).replace('-Град', '-град')


def norm(name):
    """Name compared across sources: no quotes, spaces, hyphens or case."""
    return re.sub(r'[\s„“"”\'-]', '', fix_cyrillic(name)).lower()


def soup_of(year):
    path = os.path.join(CACHE, f'zdb-{year}', f'zdb-{year}.html')
    return BeautifulSoup(open(path, encoding='utf-8').read(), 'lxml')


def caption(soup, start):
    tag = soup.find(lambda t: t.name == 'p' and clean(t.get_text()).startswith(start))
    if tag is None:
        raise ValueError(f'not found: {start}')
    return tag


def municipal_rows(table, width):
    """(province, municipality, cells) for each municipality row of a per-municipality table, and the "ВСИЧКО:" cells."""
    rows, total, province = [], None, None
    for tr in table.find_all('tr'):
        cells = [clean(td.get_text()) for td in tr.find_all(['td', 'th'])]
        if len(cells) != width or not cells[0]:
            continue
        first = fix_cyrillic(cells[0])
        if re.match(r'^ОБ\s*ЛАСТ\s', first, re.I) and not any(cells[1:]):
            province = province_of(first)
            continue
        if first.startswith('ВСИЧКО'):
            total = [amount(c) for c in cells[1:]]
            continue
        if not re.search(r'\d', cells[1]) or re.fullmatch(r'\d+', cells[0]):
            continue  # the column-number row
        name = SOFIA if first.upper() == SOFIA.upper() else first
        # Sofia follows the Smolyan rows without a heading of its own.
        rows.append(('София-град' if name == SOFIA else province, name, [amount(c) for c in cells[1:]]))
    if total is None:
        raise ValueError('no "ВСИЧКО:" row')
    return rows, total


def transfers(year):
    article, delegated_article, unit = YEARS[year]
    soup = soup_of(year)
    head = caption(soup, f'Чл. {article}. Приема размерите на бюджетните взаимоотношения между централния бюджет и бюджетите на общините')
    # The article's text up to its table: the amounts of each type.
    text, node = [], head
    table = head.find_next('table')
    while node is not None and node is not table:
        if node.name == 'p':
            text.append(clean(node.get_text()))
        node = node.find_next(['p', 'table'])
    text = ' '.join(text)
    printed = {}
    for key, _, words in TRANSFERS:
        m = re.search(words + r'\s+([\d ]+,\d)\s*хил', text)
        if not m:
            raise ValueError(f'{year}: no amount for {key} in the text of art. {article}')
        printed[key] = amount(m.group(1))

    rows, total = municipal_rows(table, 7)
    if len(rows) != 265:
        raise ValueError(f'{year}: {len(rows)} municipalities in art. {article}')
    for province, name, cells in rows:
        main = sum(c or 0 for c in cells[1:5])
        if r1(cells[0] or 0) != r1(main):
            raise ValueError(f'{year} {name}: column 2 {cells[0]} != 3+4+5+6 {main}')
    for key, column, _ in TRANSFERS:
        i = column - 2
        col_sum = sum(c[i] or 0 for _, _, c in rows)
        if not (r1(col_sum) == r1(total[i]) == r1(printed[key])):
            raise ValueError(f'{year} {key}: column {col_sum}, "ВСИЧКО:" {total[i]}, text {printed[key]}')
    if r1(total[0]) != r1(sum(total[1:5])):
        raise ValueError(f'{year}: "ВСИЧКО:" column 2 != 3+4+5+6')

    # Art. 1(2): "Предоставени трансфери за: … Общините".
    cap = caption(soup, f'(2) Приема държавния бюджет за {year} г. по разходите')
    art1 = None
    for tr in cap.find_next('table').find_all('tr'):
        cells = [clean(td.get_text()) for td in tr.find_all(['td', 'th'])]
        if len(cells) == 3 and cells[1] == 'Общините':
            art1 = amount(cells[2])
    grand = total[0] + total[5]
    if art1 is None or r1(art1) != r1(grand):
        raise ValueError(f'{year}: art. 1(2) "Общините" {art1} != {grand}')

    check_delegated(year, soup, delegated_article, rows)
    print(f'{year}: art. {article}, {len(rows)} municipalities; ' + ', '.join(f'{k} {total[c - 2]:,.1f}' for k, c, _ in TRANSFERS)
          + f'; all {grand:,.1f} = art. 1(2) "Общините" {art1:,.1f} {unit}')
    return rows, total, unit, article


def check_delegated(year, soup, article, transfer_rows):
    """Art. 52/54 by function: equals the existing CSV, and its totals equal the transfers table's column 3."""
    head = caption(soup, f'Чл. {article}. Определя стойностните показатели на общините за делегираните от държавата дейности')
    rows, total = municipal_rows(head.find_next('table'), 11)
    if [(p, n) for p, n, _ in rows] != [(p, n) for p, n, _ in transfer_rows]:
        raise ValueError(f'{year}: art. {article} lists different municipalities')
    for (_, name, cells), (_, _, tcells) in zip(rows, transfer_rows):
        if r1(cells[0] or 0) != r1(tcells[1] or 0):
            raise ValueError(f'{year} {name}: art. {article} total {cells[0]} != delegated subsidy {tcells[1]}')
        parts = sum(c or 0 for c in (cells[1], *cells[4:]))
        if r1(cells[0]) != r1(parts) or r1(cells[1] or 0) != r1((cells[2] or 0) + (cells[3] or 0)):
            raise ValueError(f'{year} {name}: art. {article} row does not add up')
    path = os.path.join(SOURCES, f'budget-{year}', f'municipal-delegated-{year}.csv')
    with open(path, newline='', encoding='utf-8') as f:
        existing = [r for r in csv.DictReader(f)]
    body = [r for r in existing if not r['municipality'].strip().upper().startswith('ВСИЧКО')]
    if len(body) != len(rows):
        raise ValueError(f'{path}: {len(body)} rows, the act {len(rows)}')
    as_float = lambda v: float(v) if v.strip() else None
    for r, (_, name, cells) in zip(body, rows):
        if norm(r['municipality']) != norm(name):
            raise ValueError(f'{path}: {r["municipality"]} != {name}')
        if [as_float(r[c]) for c in DELEGATED_COLUMNS] != cells:
            raise ValueError(f'{path}: {name} differs from the act')
    last = [r for r in existing if r['municipality'].strip().upper().startswith('ВСИЧКО')][0]
    if [as_float(last[c]) for c in DELEGATED_COLUMNS] != total:
        raise ValueError(f'{path}: "ВСИЧКО:" differs from the act')
    print(f'{year}: art. {article} (delegated activities by function) equals {os.path.relpath(path, ROOT)} in all '
          f'{len(rows)} rows and the total ({total[0]:,.1f}); every municipality\'s total equals its delegated subsidy in the transfers table')


# ---------- register: ЕБК codes and residents ----------

def ebk_codes():
    """ЕБК 2026, section VII В): municipality → code, checked against annex pr.7 of the 2025 report."""
    wb = openpyxl.load_workbook(os.path.join(CACHE, 'municipalities', 'EBK-2026-public.xlsx'), read_only=True, data_only=True)
    codes = []
    for name, code in ((r[1], r[2]) for r in wb['Local Government'].iter_rows(values_only=True)):
        if isinstance(code, (int, float)) and isinstance(name, str) and not name.strip().startswith('Район'):
            codes.append((int(code), name.strip()))
    wb = openpyxl.load_workbook(os.path.join(CACHE, 'municipalities', 'report-2025-pr7.xlsx'), read_only=True, data_only=True)
    pr7 = {}
    for row in wb.worksheets[0].iter_rows(min_row=5, values_only=True):
        if row[0] is not None:
            m = re.match(r'^Община (.+?), област', str(row[3]).strip()) or re.match(r'^(Столична община)', str(row[3]).strip())
            pr7[int(row[0])] = m.group(1)
    return codes, pr7


def match_codes(municipalities):
    codes, pr7 = ebk_codes()
    # Province code prefixes from the ЕБК's order (51 Благоевград … 78 Ямбол; 72 is Sofia).
    prefix = {}
    for province, name in municipalities:
        cands = [c for c, n in codes if norm(n) == norm(name)]
        if len(cands) == 1:
            prefix.setdefault(province, set()).add(cands[0] // 100)
    prefix = {p: s.pop() for p, s in prefix.items() if len(s) == 1}
    result = {}
    for province, name in municipalities:
        cands = [c for c, n in codes if norm(n) == norm(name) and c // 100 == prefix[province]]
        if len(cands) != 1:
            raise ValueError(f'ЕБК code of {name} ({province}): {cands}')
        code = cands[0]
        if code in pr7 and norm(pr7[code]) != norm(name):
            raise ValueError(f'ЕБК {code}: {name} in the ЕБК, {pr7[code]} in annex pr.7')
        result[(province, name)] = code
    missing = [n for (p, n), c in result.items() if c not in pr7]
    if len(set(result.values())) != len(result):
        raise ValueError('duplicate ЕБК codes')
    print(f'ЕБК codes: {len(result)} municipalities; annex pr.7 agrees on {len(result) - len(missing)} (not in pr.7: {", ".join(missing)})')
    return result


NSI_PROVINCE = {'София (столица)': 'София-град', 'София': 'Софийска'}


def nsi_residents(municipalities):
    """Residents at 31 Dec of each year: (province, municipality) → {year: residents}."""
    wb = openpyxl.load_workbook(os.path.join(CACHE, 'municipalities', 'Pop_6.1.1_Pop_DR.xlsx'), read_only=True, data_only=True)
    by_province = {}
    for province, name in municipalities:
        by_province.setdefault(province, []).append(name)
    result = {key: {} for key in municipalities}
    for year, sheet in NSI_SHEETS.items():
        rows = [(clean(str(r[0])), r[1]) for r in wb[sheet].iter_rows(values_only=True) if r[0] and isinstance(r[1], int)]
        start = next(i for i, (n, _) in enumerate(rows) if n == 'Общо за страната')
        national = rows[start][1]
        i, matched = start + 1, 0
        while i < len(rows):
            pname, ptotal = rows[i]
            province = NSI_PROVINCE.get(pname, pname)
            names = by_province.get(province)
            if names is None:
                raise ValueError(f'NSI {year}: unknown province {pname}')
            block = rows[i + 1:i + 1 + len(names)]
            if sum(v for _, v in block) != ptotal:
                raise ValueError(f'NSI {year}: {pname} municipalities do not add up')
            dobrich_city = any(n in ('Добрич - град',) for n, _ in block)
            for n, v in block:
                key = {'столична': norm(SOFIA), 'добрич-град': 'добрич', 'добрич-селска': 'добричка'}.get(n.lower().replace(' ', ''), norm(n))
                if province == 'Добрич' and n == 'Добрич' and dobrich_city:
                    key = 'добричка'  # 2023-2024 call the rural municipality "Добрич" and the town "Добрич - град"
                found = [m for m in names if norm(m) == key]
                if len(found) != 1:
                    raise ValueError(f'NSI {year}: {n} ({pname}) not matched')
                result[(province, found[0])][year] = v
                matched += 1
            i += 1 + len(names)
        if matched != 265 or sum(r[year] for r in result.values()) != national:
            raise ValueError(f'NSI {year}: {matched} municipalities')
        print(f'NSI 31.12.{year}: 265 municipalities, {national:,} residents (= the national total)')
    # The town of Dobrich is several times larger than the rural municipality around it.
    for year in NSI_SHEETS:
        if not result[('Добрич', 'Добрич')][year] > 3 * result[('Добрич', 'Добричка')][year]:
            raise ValueError(f'NSI {year}: Dobrich town and rural municipality swapped')
    return result


# ---------- output ----------

def write_transfers(year, rows, total, unit, codes):
    path = os.path.join(SOURCES, f'budget-{year}', f'municipal-transfers-{year}.csv')
    value_col = f'amount_{unit}'
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['province', 'municipality', 'ebk_code', 'transfer', 'column', value_col])
        fmt = lambda v: '' if v is None else f'{v:.1f}'
        for province, name, cells in rows:
            for key, column, _ in TRANSFERS:
                w.writerow([province, name, codes[(province, name)], key, column, fmt(cells[column - 2])])
        for key, column, _ in TRANSFERS:
            w.writerow(['', 'ВСИЧКО:', '', key, column, fmt(total[column - 2])])
    print(f'  → {os.path.relpath(path, ROOT)}')


def write_register(municipalities, codes, residents):
    os.makedirs(os.path.join(SOURCES, 'places'), exist_ok=True)
    path = os.path.join(SOURCES, 'places', 'municipalities.csv')
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['ebk_code', 'province', 'municipality'] + [f'residents_{y}_12_31' for y in NSI_SHEETS])
        for key in sorted(municipalities, key=lambda k: codes[k]):
            w.writerow([codes[key], *key] + [residents[key][y] for y in NSI_SHEETS])
    print(f'  → {os.path.relpath(path, ROOT)}')


def main():
    results = {year: transfers(year) for year in YEARS}
    lists = {year: [(p, n) for p, n, _ in r[0]] for year, r in results.items()}
    if not (lists[2024] == lists[2025] == lists[2026]):
        raise ValueError('the acts list different municipalities')
    municipalities = lists[2026]
    codes = match_codes(municipalities)
    residents = nsi_residents(municipalities)
    for year, (rows, total, unit, _) in results.items():
        write_transfers(year, rows, total, unit, codes)
    write_register(municipalities, codes, residents)


if __name__ == '__main__':
    sys.exit(main())
